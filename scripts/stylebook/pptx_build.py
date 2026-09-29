"""把一份 PPT 配图计划的出图组装成 .pptx（需要 python-pptx；没装时命令行用 uv 临时带上，不往系统 Python 里装）。

两种模式：
- full（整页图）：每页一张全幅图，文字已由模型画在图里；
- illustration（插画 + 可编辑文字）：出图时左侧 40% 留空（text.mode = overlay、reserve = the left 40% of the frame），
  这里在留空处放可编辑的标题与要点文本框，改字不用重出图。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

SLIDE_W, SLIDE_H = 12192000, 6858000  # 16:9，EMU
FONT = "PingFang SC"


def slide_text(item: dict) -> tuple[str, list[str]]:
    items = (item.get("text") or {}).get("items") or []
    title = next((x["text"] for x in items if x.get("role") == "title"), item.get("what", ""))
    bullets = item.get("slide_points") or item.get("points") or []
    return title, bullets


def build(plan: dict, images: Path, out: Path, mode: str = "illustration") -> Path:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    margin = int(SLIDE_W * 0.05)
    for it in plan["items"]:
        img = Path(images) / f"{it['id']}.png"
        if not img.is_file():
            raise FileNotFoundError(img)
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(img), 0, 0, width=Emu(SLIDE_W), height=Emu(SLIDE_H))
        title, bullets = slide_text(it)
        s.notes_slide.notes_text_frame.text = it.get("why", "")
        if mode == "full":
            continue
        box_w = int(SLIDE_W * 0.40) - margin
        tb = s.shapes.add_textbox(Emu(margin), Emu(int(SLIDE_H * 0.12)), Emu(box_w), Emu(int(SLIDE_H * 0.18)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = title
        for r in tf.paragraphs[0].runs:
            r.font.size, r.font.bold, r.font.name = Pt(36), True, FONT
        if bullets:
            bb = s.shapes.add_textbox(Emu(margin), Emu(int(SLIDE_H * 0.34)), Emu(box_w), Emu(int(SLIDE_H * 0.55)))
            bf = bb.text_frame
            bf.word_wrap = True
            for i, b in enumerate(bullets):
                p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
                p.text = f"• {b}"
                p.space_after = Pt(10)
                for r in p.runs:
                    r.font.size, r.font.name = Pt(20), FONT
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":  # uv run --with python-pptx python3 pptx_build.py <plan> <images> <out> <mode>
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    plan_p, img_d, out_p, mode = sys.argv[1:5]
    print(build(json.loads(Path(plan_p).read_text(encoding="utf-8")), Path(img_d), Path(out_p), mode))
